"""Lion wafer-level Good Die report cleaner.

This module handles the approved ``WAFER PROBE FAIL COUNTER REPORT`` and
LCD235 workbooks used for monthly die-count consolidation.  It is
intentionally separate from the existing Lion CP die-level V1/V2 pipelines
because its output is a five-column business summary rather than the standard
cleaned/yield/spec CSV contract.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from cp_data_processor.processing.output_naming import create_output_run_dir


OUTPUT_COLUMNS = ("NCE品名", "LOT", "Wafer", "PASS", "Good Die")
REPORT_MARKER = "WAFER PROBE FAIL COUNTER REPORT"
LCD235_TITLE_MARKER = "杭州立昂微电子MOSFET出货报告单"
LCD235_APPROVED_HEADERS = (
    "",
    "型号",
    "批号",
    "片号",
    "测试良率",
    "CP合格管芯数",
    "QAD补点数",
    "QAD良率",
    "IGSSF_5",
    "IGSSF_20",
    "IGSSR_20",
    "IGSSF_30",
    "IGSSR_30",
    "IGSSF_40",
    "IGSSR_40",
    "VTH1",
    "IDSS_1500",
    "HBVDSS1-100uA",
    "HBVDSS2-250uA",
    "HBVDSS3-1mA",
    "HBVDSS4-50uA",
    "HBVDSS5-250uA",
    "VTH2",
    "DVR_BVDSS2",
    "DVR_BVDSS3",
    "DVR_BVDSS4",
    "DVR_BVDSS5",
    "RDSON",
    "VFSD",
    "IDSS_1200",
    "IDSS_50",
    "IDSS_1500_Retest",
    "IGSSF_20_Retest",
    "IGSSR_20_Retest",
)
LCD235_SUMMARY_HEADERS = (
    "总片数",
    "CP合格管芯总数",
    "合格管芯总数",
    "QAD平均良率",
)


class LionDieCountError(ValueError):
    """Raised when a source workbook does not match the approved format."""


@dataclass(frozen=True)
class LionDieCountRecord:
    product: str
    lot_id: str
    wafer_id: int
    pass_count: int
    good_die: int
    source_file: Path


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalized_text(value: object) -> str:
    return "".join(_text(value).split())


def _filename_matches_lot(source: Path, lot_id: str) -> bool:
    """Allow the real lot ID followed by a space-separated user note.

    The workbook metadata remains the identity source.  A note must begin
    after whitespace, so a filename with an unrelated or merely similar lot
    ID is still rejected.
    """

    if source.stem == lot_id:
        return True
    suffix = source.stem.removeprefix(lot_id)
    return bool(suffix) and suffix[0].isspace() and bool(suffix.strip())


def discover_workbooks(input_dir: str | Path) -> tuple[list[Path], list[Path]]:
    """Recursively find source workbooks and separate Excel lock files."""

    root = Path(input_dir).expanduser()
    if not root.exists():
        raise LionDieCountError(f"输入目录不存在: {root}")
    if not root.is_dir():
        raise LionDieCountError(f"输入必须是目录: {root}")

    all_xlsx = sorted(
        root.rglob("*.xlsx"), key=lambda path: str(path.relative_to(root)).casefold()
    )
    temporary = [path for path in all_xlsx if path.name.startswith("~$")]
    workbooks = [path for path in all_xlsx if not path.name.startswith("~$")]
    if not workbooks:
        raise LionDieCountError(f"目录中未找到有效的 Lion 管芯数 .xlsx 文件: {root}")
    return workbooks, temporary


def _extract_identity(metadata_row: Iterable[object], source_file: Path) -> tuple[str, str]:
    values = [_text(value) for value in metadata_row if _text(value)]
    device_values = [
        value.split("=", 1)[1].strip()
        for value in values
        if value.upper().startswith("DEVICE=")
    ]
    lot_values = [
        value.split("=", 1)[1].strip()
        for value in values
        if value.upper().startswith("LOT#=")
    ]
    if len(device_values) != 1 or not device_values[0]:
        raise LionDieCountError(
            f"{source_file.name}: 第2行必须且只能包含一个有效 DEVICE="
        )
    if len(lot_values) != 1 or not lot_values[0]:
        raise LionDieCountError(
            f"{source_file.name}: 第2行必须且只能包含一个有效 LOT#="
        )
    return device_values[0], lot_values[0]


def _as_non_negative_integer(value: object, field: str, source_file: Path, row_no: int) -> int:
    try:
        number = float(value)
        integer = int(number)
    except (TypeError, ValueError, OverflowError) as exc:
        raise LionDieCountError(
            f"{source_file.name}: 第{row_no}行 {field} 不是整数: {value!r}"
        ) from exc
    if number != integer or integer < 0:
        raise LionDieCountError(
            f"{source_file.name}: 第{row_no}行 {field} 必须是非负整数: {value!r}"
        )
    return integer


def _read_probe_counter_rows(
    source: Path, rows: list[tuple[object, ...]]
) -> list[LionDieCountRecord]:
    """Read the approved ``WAFER PROBE FAIL COUNTER REPORT`` format."""

    if len(rows) < 4:
        raise LionDieCountError(f"{source.name}: 文件行数不足")
    product, lot_id = _extract_identity(rows[1], source)
    if not _filename_matches_lot(source, lot_id):
        raise LionDieCountError(
            f"{source.name}: LOT#={lot_id} 与文件名不一致；仅允许“批号”或“批号 + 空格 + 备注”"
        )

    if not any(REPORT_MARKER in _text(value).upper() for row in rows for value in row):
        raise LionDieCountError(f"{source.name}: 未找到 {REPORT_MARKER} 格式标识")

    header_matches: list[tuple[int, list[str]]] = []
    for index, row in enumerate(rows):
        normalized = [_text(value).upper() for value in row]
        if "WAFER#" in normalized and "PASS" in normalized and "DIE" in normalized:
            header_matches.append((index, normalized))
    if len(header_matches) != 1:
        raise LionDieCountError(
            f"{source.name}: 应唯一找到包含 WAFER#、PASS 和 DIE 的表头，实际 {len(header_matches)} 个"
        )

    header_index, header = header_matches[0]
    wafer_index = header.index("WAFER#")
    pass_index = header.index("PASS")
    die_index = header.index("DIE")
    records: list[LionDieCountRecord] = []
    summary_index: int | None = None

    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        wafer_value = row[wafer_index] if wafer_index < len(row) else None
        wafer_text = _text(wafer_value)
        if wafer_text.upper() == "SUMMARY":
            summary_index = row_index - 1
            break
        if not wafer_text or set(wafer_text) <= {"-"}:
            continue
        pass_value = row[pass_index] if pass_index < len(row) else None
        die_value = row[die_index] if die_index < len(row) else None
        wafer_id = _as_non_negative_integer(wafer_value, "WAFER#", source, row_index)
        pass_count = _as_non_negative_integer(pass_value, "PASS", source, row_index)
        good_die = _as_non_negative_integer(die_value, "DIE", source, row_index)
        if wafer_id <= 0:
            raise LionDieCountError(
                f"{source.name}: 第{row_index}行 WAFER# 必须大于0"
            )
        records.append(
            LionDieCountRecord(
                product, lot_id, wafer_id, pass_count, good_die, source
            )
        )

    if not records:
        raise LionDieCountError(f"{source.name}: 未找到 Wafer 数据行")
    if summary_index is None:
        raise LionDieCountError(f"{source.name}: 未找到 SUMMARY 行")

    summary_data = None
    for row_index in range(summary_index + 1, len(rows)):
        row = rows[row_index]
        first = _text(row[wafer_index] if wafer_index < len(row) else None)
        if first and not set(first) <= {"-"}:
            summary_data = (row_index + 1, row)
            break
    if summary_data is None:
        raise LionDieCountError(f"{source.name}: SUMMARY 后缺少汇总数据")

    summary_row_no, summary_row = summary_data
    summary_wafer_count = _as_non_negative_integer(
        summary_row[wafer_index], "SUMMARY WAFER#", source, summary_row_no
    )
    summary_pass = _as_non_negative_integer(
        summary_row[pass_index], "SUMMARY PASS", source, summary_row_no
    )
    summary_good_die = _as_non_negative_integer(
        summary_row[die_index], "SUMMARY DIE", source, summary_row_no
    )
    if summary_wafer_count != len(records):
        raise LionDieCountError(
            f"{source.name}: SUMMARY Wafer数={summary_wafer_count}，明细行数={len(records)}"
        )
    actual_pass = sum(record.pass_count for record in records)
    if summary_pass != actual_pass:
        raise LionDieCountError(
            f"{source.name}: SUMMARY PASS={summary_pass}，明细合计={actual_pass}"
        )
    actual_good_die = sum(record.good_die for record in records)
    if summary_good_die != actual_good_die:
        raise LionDieCountError(
            f"{source.name}: SUMMARY DIE={summary_good_die}，明细合计={actual_good_die}"
        )

    wafer_ids = [record.wafer_id for record in records]
    if len(wafer_ids) != len(set(wafer_ids)):
        raise LionDieCountError(f"{source.name}: 存在重复 WAFER#")
    return records


def _read_lcd235_rows(
    source: Path, rows: list[tuple[object, ...]]
) -> list[LionDieCountRecord]:
    """Read the approved LCD235 report format and reconcile its summary."""

    if len(rows) < 10:
        raise LionDieCountError(f"{source.name}: LCD235 文件行数不足")
    if not any(
        LCD235_TITLE_MARKER in _normalized_text(value)
        for row in rows[:5]
        for value in row
    ):
        raise LionDieCountError(f"{source.name}: 未找到 LCD235 报告标识")

    header_matches: list[int] = []
    for index, row in enumerate(rows):
        normalized = tuple(_normalized_text(value) for value in row)
        if normalized == LCD235_APPROVED_HEADERS:
            header_matches.append(index)
    if len(header_matches) != 1:
        raise LionDieCountError(
            f"{source.name}: LCD235 34列表头应唯一匹配已批准顺序，实际 {len(header_matches)} 个"
        )

    header_index = header_matches[0]
    header = LCD235_APPROVED_HEADERS
    product_index = header.index("型号")
    lot_index = header.index("批号")
    wafer_index = header.index("片号")
    pass_index = header.index("CP合格管芯数")
    qad_index = header.index("QAD补点数")

    summary_matches = [
        index
        for index, row in enumerate(rows[header_index + 1 :], start=header_index + 1)
        if _normalized_text(row[product_index] if product_index < len(row) else None)
        == "摘要"
    ]
    if len(summary_matches) != 1:
        raise LionDieCountError(
            f"{source.name}: LCD235 应唯一找到摘要行，实际 {len(summary_matches)} 个"
        )
    summary_index = summary_matches[0]
    if summary_index + 2 >= len(rows):
        raise LionDieCountError(f"{source.name}: LCD235 摘要后缺少汇总数据")

    summary_header = rows[summary_index + 1]
    normalized_summary = tuple(
        _normalized_text(value) for value in summary_header[4:8]
    )
    if normalized_summary != LCD235_SUMMARY_HEADERS:
        raise LionDieCountError(
            f"{source.name}: LCD235 汇总表头与已批准格式不一致"
        )

    records: list[LionDieCountRecord] = []
    expected_product: str | None = None
    expected_lot: str | None = None
    data_start_index = header_index + 3
    for zero_based_index, row in enumerate(
        rows[data_start_index:summary_index], start=data_start_index
    ):
        row_no = zero_based_index + 1
        if not any(_text(value) for value in row):
            continue
        product = _text(row[product_index] if product_index < len(row) else None)
        lot_id = _text(row[lot_index] if lot_index < len(row) else None)
        if not product or not lot_id:
            raise LionDieCountError(
                f"{source.name}: 第{row_no}行缺少型号或批号"
            )
        if not _filename_matches_lot(source, lot_id):
            raise LionDieCountError(
                f"{source.name}: 第{row_no}行批号={lot_id} 与文件名不一致；仅允许“批号”或“批号 + 空格 + 备注”"
            )
        if expected_product is None:
            expected_product = product
            expected_lot = lot_id
        if product != expected_product or lot_id != expected_lot:
            raise LionDieCountError(
                f"{source.name}: 第{row_no}行型号/批号与本文件其他 Wafer 不一致"
            )

        wafer_id = _as_non_negative_integer(
            row[wafer_index] if wafer_index < len(row) else None,
            "片号",
            source,
            row_no,
        )
        pass_count = _as_non_negative_integer(
            row[pass_index] if pass_index < len(row) else None,
            "CP合格管芯数",
            source,
            row_no,
        )
        qad_count = _as_non_negative_integer(
            row[qad_index] if qad_index < len(row) else None,
            "QAD补点数",
            source,
            row_no,
        )
        if wafer_id <= 0:
            raise LionDieCountError(f"{source.name}: 第{row_no}行片号必须大于0")
        good_die = pass_count - qad_count
        if good_die < 0:
            raise LionDieCountError(
                f"{source.name}: 第{row_no}行 QAD补点数={qad_count} 大于 CP合格管芯数={pass_count}"
            )
        records.append(
            LionDieCountRecord(
                product, lot_id, wafer_id, pass_count, good_die, source
            )
        )

    if not records:
        raise LionDieCountError(f"{source.name}: LCD235 未找到 Wafer 数据行")
    wafer_ids = [record.wafer_id for record in records]
    if len(wafer_ids) != len(set(wafer_ids)):
        raise LionDieCountError(f"{source.name}: LCD235 存在重复片号")

    summary_row_no = summary_index + 3
    summary_row = rows[summary_index + 2]
    summary_wafer_count = _as_non_negative_integer(
        summary_row[4], "汇总总片数", source, summary_row_no
    )
    summary_pass = _as_non_negative_integer(
        summary_row[5], "汇总CP合格管芯总数", source, summary_row_no
    )
    summary_good_die = _as_non_negative_integer(
        summary_row[6], "汇总合格管芯总数", source, summary_row_no
    )
    actual_pass = sum(record.pass_count for record in records)
    actual_good_die = sum(record.good_die for record in records)
    if summary_wafer_count != len(records):
        raise LionDieCountError(
            f"{source.name}: 汇总总片数={summary_wafer_count}，明细行数={len(records)}"
        )
    if summary_pass != actual_pass:
        raise LionDieCountError(
            f"{source.name}: 汇总CP合格管芯总数={summary_pass}，明细合计={actual_pass}"
        )
    if summary_good_die != actual_good_die:
        raise LionDieCountError(
            f"{source.name}: 汇总合格管芯总数={summary_good_die}，明细计算={actual_good_die}"
        )
    return records


def read_workbook(source_file: str | Path) -> list[LionDieCountRecord]:
    """Read one approved Lion die-count workbook and dispatch by exact format."""

    source = Path(source_file)
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise LionDieCountError(f"无法读取 {source.name}: {exc}") from exc

    try:
        sheet_names = workbook.sheetnames
        if sheet_names == ["Sheet"]:
            rows = list(workbook["Sheet"].iter_rows(values_only=True))
            return _read_probe_counter_rows(source, rows)
        if sheet_names == ["Sheet1"]:
            rows = list(workbook["Sheet1"].iter_rows(values_only=True))
            return _read_lcd235_rows(source, rows)
        raise LionDieCountError(
            f"{source.name}: Sheet 结构不受支持: {sheet_names}"
        )
    finally:
        workbook.close()


def write_output(records: list[LionDieCountRecord], output_file: str | Path) -> Path:
    """Write the exact five-column business workbook shown in the requirement."""

    target = Path(output_file)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Lion管芯数"
    worksheet.append(OUTPUT_COLUMNS)
    for record in records:
        worksheet.append(
            (
                record.product,
                record.lot_id,
                record.wafer_id,
                record.pass_count,
                record.good_die,
            )
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 14
    worksheet.column_dimensions["E"].width = 14
    workbook.save(target)
    workbook.close()
    return target


def process_lion_die_count_directory(
    input_dir: str | Path,
    output_parent: str | Path,
    *,
    progress: Callable[[str], None] | None = None,
) -> dict[str, object]:
    """Clean all monthly Lion die-count workbooks into one Excel report."""

    emit = progress or (lambda _message: None)
    workbooks, temporary_files = discover_workbooks(input_dir)
    emit(f"已发现 {len(workbooks)} 个有效 .xlsx 文件。")
    if temporary_files:
        emit(f"已忽略 {len(temporary_files)} 个 Excel 临时锁文件（~$...）。")

    all_records: list[LionDieCountRecord] = []
    seen_keys: dict[tuple[str, str, int], Path] = {}
    for index, source_file in enumerate(workbooks, start=1):
        records = read_workbook(source_file)
        for record in records:
            key = (record.product, record.lot_id, record.wafer_id)
            previous = seen_keys.get(key)
            if previous is not None:
                raise LionDieCountError(
                    f"重复的 NCE品名+LOT+Wafer: {key}；来源 {previous.name} 和 {source_file.name}"
                )
            seen_keys[key] = source_file
        all_records.extend(records)
        if index == 1 or index == len(workbooks) or index % 25 == 0:
            emit(f"已校验 {index}/{len(workbooks)} 个文件。")

    first_lot_id = all_records[0].lot_id
    output_dir = create_output_run_dir(output_parent, first_lot_id)
    output_file = write_output(all_records, output_dir / "Lion_管芯数.xlsx")
    emit(f"已生成 {len(all_records)} 条 Wafer 记录。")
    return {
        "output_dir": str(output_dir),
        "output_file": str(output_file),
        "file_count": len(workbooks),
        "ignored_temp_file_count": len(temporary_files),
        "product_count": len({record.product for record in all_records}),
        "lot_count": len({record.lot_id for record in all_records}),
        "wafer_count": len(all_records),
        "first_lot_id": first_lot_id,
    }
