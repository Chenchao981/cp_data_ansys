from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from lion.lion_die_count_processor import (
    LionDieCountError,
    OUTPUT_COLUMNS,
    process_lion_die_count_directory,
    read_workbook,
)


def make_source(
    path: Path,
    *,
    product="NCELFR140EB20BA",
    lot_id="F26240105",
    rows=((9, 883), (10, 977)),
    summary_pass=None,
):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet"
    worksheet.append(["2026-06-01", None, None, None, f"{lot_id}.{product}"])
    worksheet.append([f"LOT#={lot_id}", None, None, None, "TESTFILE=F.pgs", None, None, None, f"DEVICE={product}"])
    worksheet.append(["WAFER PROBE FAIL COUNTER REPORT"])
    worksheet.append(["WAFER#", "#2", "PASS", "PASS%", "TOTAL"])
    worksheet.append(["---------------", "------", "------"])
    for wafer_id, good_die in rows:
        worksheet.append([str(wafer_id), "0", str(good_die), "99%", str(good_die)])
    worksheet.append(["SUMMARY", "#2", "PASS", "PASS%", "TOTAL"])
    worksheet.append(["---------------", "------", "------"])
    worksheet.append(
        [str(len(rows)), "0", str(sum(value for _, value in rows) if summary_pass is None else summary_pass)]
    )
    workbook.save(path)
    workbook.close()


def test_recursive_monthly_cleaning_writes_exact_contract(tmp_path):
    input_dir = tmp_path / "2026.6"
    make_source(input_dir / "F0100A1" / "F26240105.xlsx")
    make_source(
        input_dir / "F0200A1" / "F26240106.xlsx",
        product="NCELFR75EV40AA",
        lot_id="F26240106",
        rows=((1, 1000),),
    )
    (input_dir / "F0100A1" / "~$F26240105.xlsx").touch()

    messages = []
    result = process_lion_die_count_directory(
        input_dir, tmp_path / "output", progress=messages.append
    )

    assert result["file_count"] == 2
    assert result["ignored_temp_file_count"] == 1
    assert result["product_count"] == 2
    assert result["lot_count"] == 2
    assert result["wafer_count"] == 3
    assert Path(result["output_dir"]).name.startswith("F26240105_")
    workbook = load_workbook(result["output_file"], read_only=True, data_only=True)
    rows = list(workbook["Lion管芯数"].iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == OUTPUT_COLUMNS
    assert rows[1:] == [
        ("NCELFR140EB20BA", "F26240105", 9, 883),
        ("NCELFR140EB20BA", "F26240105", 10, 977),
        ("NCELFR75EV40AA", "F26240106", 1, 1000),
    ]
    assert any("已忽略 1 个 Excel 临时锁文件" in message for message in messages)


def test_summary_pass_mismatch_fails_closed(tmp_path):
    source = tmp_path / "F26240105.xlsx"
    make_source(source, summary_pass=1)

    with pytest.raises(LionDieCountError, match="SUMMARY PASS"):
        read_workbook(source)


def test_filename_must_match_lot_metadata(tmp_path):
    source = tmp_path / "WRONG.xlsx"
    make_source(source)

    with pytest.raises(LionDieCountError, match="LOT#=.*文件名不一致"):
        read_workbook(source)


def test_duplicate_product_lot_wafer_across_files_fails_closed(tmp_path):
    input_dir = tmp_path / "input"
    make_source(input_dir / "a" / "F26240105.xlsx")
    make_source(input_dir / "b" / "F26240105.xlsx")

    with pytest.raises(LionDieCountError, match=r"重复的 NCE品名\+LOT\+Wafer"):
        process_lion_die_count_directory(input_dir, tmp_path / "output")
