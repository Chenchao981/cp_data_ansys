from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from cp_data_processor.analysis.quick_pat import (
    FORMULA_CONTRACT,
    generate_cleaned_csv_pat,
)


def test_generate_quick_pat_uses_vdmos_position_quantiles_and_spec_parameters(
    tmp_path: Path,
) -> None:
    cleaned = tmp_path / "LOT_cleaned_20260902.csv"
    spec = tmp_path / "LOT_spec_20260902.csv"
    pd.DataFrame(
        {
            "Lot_ID": ["LOT"] * 12,
            "Wafer_ID": ["1"] * 12,
            "X": list(range(12)),
            "Y": list(range(12)),
            "Bin": [1] * 12,
            "VTH": list(range(12)),
            "CONT": [99] * 12,
        }
    ).to_csv(cleaned, index=False)
    pd.DataFrame(
        {"Parameter": ["VTH"], "Unit": ["V"], "LimitL": [0], "LimitU": [10]}
    ).to_csv(spec, index=False)

    report = generate_cleaned_csv_pat(
        cleaned_files=[cleaned], spec_files=[spec], output_dir=tmp_path / "result"
    )

    workbook = load_workbook(report, read_only=True, data_only=True)
    try:
        sheet = workbook["PAT"]
        assert sheet["B1"].value == FORMULA_CONTRACT
        headers = [cell.value for cell in sheet[2]]
        values = [cell.value for cell in sheet[3]]
        row = dict(zip(headers, values, strict=True))
        assert row["parameter"] == "VTH"
        assert row["count"] == 12
        assert row["q1"] == 3
        assert row["median"] == 6
        assert row["q3"] == 9
        assert row["sigma"] == pytest.approx(6 / 1.349)
        assert "CONT" not in [sheet.cell(index, 1).value for index in range(3, sheet.max_row + 1)]
    finally:
        workbook.close()


def test_generate_quick_pat_reads_matrix_specs_and_rejects_short_parameters(
    tmp_path: Path,
) -> None:
    cleaned = tmp_path / "LOT_cleaned.csv"
    spec = tmp_path / "LOT_spec.csv"
    pd.DataFrame({"Lot_ID": ["LOT"] * 9, "VTH": list(range(9))}).to_csv(
        cleaned, index=False
    )
    pd.DataFrame(
        {"Parameter": ["Unit", "LimitL", "LimitU"], "VTH": ["V", "0", "10"]}
    ).to_csv(spec, index=False)

    with pytest.raises(ValueError, match="at least 10"):
        generate_cleaned_csv_pat(
            cleaned_files=[cleaned], spec_files=[spec], output_dir=tmp_path / "result"
        )


def test_generate_quick_pat_skips_declared_non_numeric_parameters(tmp_path: Path) -> None:
    cleaned = tmp_path / "LOT_cleaned.csv"
    spec = tmp_path / "LOT_spec.csv"
    pd.DataFrame(
        {"Lot_ID": ["LOT"] * 10, "TEST_NUM": [""] * 10, "VTH": list(range(10))}
    ).to_csv(cleaned, index=False)
    pd.DataFrame(
        {
            "Parameter": ["Unit", "LimitL", "LimitU"],
            "TEST_NUM": ["", "", ""],
            "VTH": ["V", "0", "10"],
        }
    ).to_csv(spec, index=False)

    report = generate_cleaned_csv_pat(
        cleaned_files=[cleaned], spec_files=[spec], output_dir=tmp_path / "result"
    )
    workbook = load_workbook(report, read_only=True, data_only=True)
    try:
        assert workbook["PAT"].cell(3, 1).value == "VTH"
        assert workbook["PAT"].max_row == 3
    finally:
        workbook.close()
